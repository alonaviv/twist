import os
from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from song_signup.models import TicketOrder, SING_SKU, ATTN_SKU
from song_signup.views import _process_tickchak_orders
from openpyxl import load_workbook
from constance import config

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

EVENT_SKU = 78079
EVENT_DATE = "15.3.25"
EVENT_NAME = "Open Mic - Babu Bar - 15.3.25"

class TickchakOrdersTest(TestCase):
    def _assert_expected_orders(self, expected_orders):
        for expected in expected_orders:
            self.assertTrue(
                TicketOrder.objects.filter(
                    order_id=expected["order_id"],
                    event_sku=EVENT_SKU,
                    ticket_type=expected["ticket_type"],
                    event_name=EVENT_NAME,
                    num_tickets=expected["num_tickets"],
                    customer_name=expected["customer_name"],
                    phone_number=expected["phone_number"]
                ).exists()
            )

    def test_process_tickchak_orders(self):
        expected_orders = [
            dict(order_id=10869770, num_tickets=2, ticket_type=SING_SKU, customer_name="Jonathan Griffin",
                 is_freebie=False, logged_in_customers=[], phone_number="059-9486359"),
            dict(order_id=10869772, num_tickets=2, ticket_type=SING_SKU, customer_name="בת שבע דיין", is_freebie=False,
                 logged_in_customers=[], phone_number="056-7069207"),
            dict(order_id=10869772, num_tickets=1, ticket_type=ATTN_SKU, customer_name="בת שבע דיין", is_freebie=False,
                 logged_in_customers=[], phone_number="056-7069207"),
            dict(order_id=10869778, num_tickets=2, ticket_type=SING_SKU, customer_name="דוד מזרחי", is_freebie=False,
                 logged_in_customers=[], phone_number="053-9509035"),
            dict(order_id=10869779, num_tickets=2, ticket_type=SING_SKU, customer_name="שי-לי שרון", is_freebie=False,
                 logged_in_customers=[], phone_number="055-6315800"),
            dict(order_id=10869781, num_tickets=1, ticket_type=SING_SKU, customer_name="Victoria Austin",
                 is_freebie=False, logged_in_customers=[], phone_number="054-5856614"),
            dict(order_id=10869788, num_tickets=1, ticket_type=SING_SKU, customer_name="חיה מזרחי", is_freebie=False,
                 logged_in_customers=[], phone_number="055-6553223"),
            dict(order_id=10869788, num_tickets=2, ticket_type=ATTN_SKU, customer_name="חיה מזרחי", is_freebie=False,
                 logged_in_customers=[], phone_number="055-6553223"),
            dict(order_id=10869792, num_tickets=1, ticket_type=SING_SKU, customer_name="אלה אטיאס", is_freebie=False,
                 logged_in_customers=[], phone_number="055-6286566"),
            dict(order_id=10869792, num_tickets=1, ticket_type=ATTN_SKU, customer_name="אלה אטיאס", is_freebie=False,
                 logged_in_customers=[], phone_number="055-6286566"),
            dict(order_id=10921071, num_tickets=3, ticket_type=SING_SKU, customer_name="איילה מחאמיד", is_freebie=False,
                 logged_in_customers=[], phone_number="058-3381903"),
            dict(order_id=11017601, num_tickets=1, ticket_type=SING_SKU, customer_name="ישי מטר", is_freebie=False,
                 logged_in_customers=[], phone_number="057-2955854"),
            dict(order_id=11017601, num_tickets=1, ticket_type=ATTN_SKU, customer_name="ישי מטר", is_freebie=False,
                 logged_in_customers=[], phone_number="057-2955854"),
            dict(order_id=11017747, num_tickets=1, ticket_type=SING_SKU, customer_name="נטע גבאי", is_freebie=False,
                 logged_in_customers=[], phone_number="051-6606754"),
        ]

        with open(os.path.join(BASE_DIR, "tickchak-orders.xlsx"), "rb") as f:
            spreadsheet_file = SimpleUploadedFile("test.xlsx", f.read())

        result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE, True)
        self.assertEqual(TicketOrder.objects.count(), 14)

        self._assert_expected_orders(expected_orders)

        self.assertEqual(result['num_orders'], 10)
        self.assertEqual(result['num_ticket_orders'], 14)
        self.assertEqual(result['num_new_ticket_orders'], 14)
        self.assertEqual(result['num_existing_ticket_orders'], 0)
        self.assertEqual(result['num_audience_tickets'], 5)
        self.assertEqual(result['num_singing_tickets'], 16)
        self.assertEqual(result['total_tickets'], 21)
        self.assertEqual(result['event_sku'], EVENT_SKU)
        self.assertEqual(result['event_name'], EVENT_NAME)

        generated_cheat_code = result['cheat_code']
        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test with a spreadsheet that has a few more orders - see that I handle duplicates correctly
        expected_orders.extend([
            dict(order_id=11024153, num_tickets=3, ticket_type=SING_SKU, customer_name="Jacqueline Gutierrez",
                 is_freebie=False, logged_in_customers=[], phone_number="058-7764017"),
            dict(order_id=11029561, num_tickets=1, ticket_type=SING_SKU, customer_name="Angela Fischer",
                 is_freebie=False, logged_in_customers=[], phone_number="050-9210003"),
            dict(order_id=11029376, num_tickets=2, ticket_type=ATTN_SKU, customer_name="מיכל מנחם", is_freebie=False,
                 logged_in_customers=[], phone_number="055-1615239"),
            dict(order_id=11029866, num_tickets=2, ticket_type=SING_SKU, customer_name="Brooke Hicks", is_freebie=False,
                 logged_in_customers=[], phone_number="057-6138697"),
            dict(order_id=11029866, num_tickets=2, ticket_type=ATTN_SKU, customer_name="Brooke Hicks", is_freebie=False,
                 logged_in_customers=[], phone_number="057-6138697"),
            dict(order_id=11038489, num_tickets=4, ticket_type=SING_SKU, customer_name="דוד מזרחי",
                 is_freebie=False, logged_in_customers=[], phone_number="053-9509035"),
            dict(order_id=11032259, num_tickets=1, ticket_type=ATTN_SKU, customer_name="אלה אטיאס",
                 is_freebie=False, logged_in_customers=[], phone_number="055-6286566"),
        ])

        with open(os.path.join(BASE_DIR, "tickchak-orders-extended.xlsx"), "rb") as f:
            spreadsheet_file = SimpleUploadedFile("test2.xlsx", f.read())

        result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE, True, True)

        self.assertEqual(TicketOrder.objects.count(), 21)
        self._assert_expected_orders(expected_orders)

        self.assertEqual(result['num_orders'], 16)
        self.assertEqual(result['num_ticket_orders'], 21)
        self.assertEqual(result['num_new_ticket_orders'], 7)
        self.assertEqual(result['num_existing_ticket_orders'], 14)
        self.assertEqual(result['num_audience_tickets'], 10)
        self.assertEqual(result['num_singing_tickets'], 26)
        self.assertEqual(result['total_tickets'], 36)
        self.assertEqual(result['event_sku'], EVENT_SKU)
        self.assertEqual(result['event_name'], EVENT_NAME)

        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test loading the exact same spreadsheet again - verify no changes
        result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE, True, True)

        self.assertEqual(TicketOrder.objects.count(), 21)
        self._assert_expected_orders(expected_orders)

        self.assertEqual(result['num_orders'], 16)
        self.assertEqual(result['num_ticket_orders'], 21)
        self.assertEqual(result['num_new_ticket_orders'], 0)
        self.assertEqual(result['num_existing_ticket_orders'], 21)
        self.assertEqual(result['num_audience_tickets'], 10)
        self.assertEqual(result['num_singing_tickets'], 26)
        self.assertEqual(result['total_tickets'], 36)
        self.assertEqual(result['event_sku'], EVENT_SKU)
        self.assertEqual(result['event_name'], EVENT_NAME)

        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test loading the spreadsheet with a different name as an initial upload
        with self.assertRaises(ValueError) as e:
            result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE+'1', True, False)
        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test re-uploading with incorrect SKU
        with self.assertRaises(ValueError):
            result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU+1, EVENT_DATE, True, True)
        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test re-uploading with incorrect date
        with self.assertRaises(ValueError):
            result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE+'1', True, True)
        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

        # Test re-uploading with same SKU, without marking duplicate
        with self.assertRaises(ValueError):
            result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE, True, False)
        self.assertEqual(config.EVENT_SKU, EVENT_SKU)
        self.assertEqual(config.FREEBIE_TICKET, generated_cheat_code)

    def test_reserved_singer_tickets(self):
        """A real tickchak spreadsheet with reserved singer tickets, with personal details replaced"""
        expected_orders = [
            # Reserved singer tickets only
            dict(order_id=891121986, num_tickets=1, ticket_type=SING_SKU, customer_name="מאיה כהן",
                 phone_number="052-1872268"),
            dict(order_id=891155150, num_tickets=1, ticket_type=SING_SKU, customer_name="Liam King",
                 phone_number="050-0607804"),
            # Reserved singer ticket with audience tickets
            dict(order_id=891123331, num_tickets=1, ticket_type=SING_SKU, customer_name="ליאור פרץ",
                 phone_number="052-2367338"),
            dict(order_id=891123331, num_tickets=3, ticket_type=ATTN_SKU, customer_name="ליאור פרץ",
                 phone_number="052-2367338"),
            # Regular singer tickets
            dict(order_id=891121905, num_tickets=2, ticket_type=SING_SKU, customer_name="Ethan Walker",
                 phone_number="053-2115753"),
            # Audience tickets
            dict(order_id=891121978, num_tickets=2, ticket_type=ATTN_SKU, customer_name="עומר ביטון",
                 phone_number="055-3197138"),
        ]

        with open(os.path.join(BASE_DIR, "tickchak-orders-reserved.xlsx"), "rb") as f:
            spreadsheet_file = SimpleUploadedFile("test.xlsx", f.read())

        result = _process_tickchak_orders(spreadsheet_file, EVENT_SKU, EVENT_DATE)
        self.assertEqual(TicketOrder.objects.count(), 58)
        self._assert_expected_orders(expected_orders)

        self.assertEqual(result['num_orders'], 57)
        self.assertEqual(result['num_ticket_orders'], 58)
        self.assertEqual(result['num_singing_tickets'], 36)
        self.assertEqual(result['num_audience_tickets'], 53)

        # Reserved singers log in like regular singers, and audience as usual
        config.PASSCODE = 'dev'
        for ticket_type, first_name, last_name, order_id in [('singer', 'מאיה', 'כהן', 891121986),
                                                             ('singer', 'Ethan', 'Walker', 891121905),
                                                             ('audience', 'עומר', 'ביטון', 891121978)]:
            with self.subTest(ticket_type=ticket_type, order_id=order_id):
                response = self.client.post(reverse('login'), {'ticket-type': ticket_type, 'first-name': first_name,
                                                               'last-name': last_name, 'passcode': 'dev',
                                                               'order-id': order_id, 'no-upload': 'on'})
                self.assertEqual(response.status_code, 200)
                self.client.logout()

    def test_order_with_two_rows_of_a_ticket_type(self):
        """An order can't hold two rows of the same ticket type, so the upload fails without saving anything"""
        cases = [
            # Order 891121905 has 2 regular singer tickets; add a reserved one with the same quantity
            (891121905, 'כרטיס זמר/ת - משוריין לזוכים!', "Order 891121905 has more than one singer ticket row"),
            # Order 891121978 has 2 audience tickets; add the same audience row again
            (891121978, 'כרטיס קהל - שרים מהכיסא!', "Order 891121978 has more than one audience ticket row"),
        ]
        for order_id, ticket_desc, error in cases:
            with self.subTest(ticket_desc=ticket_desc):
                workbook = load_workbook(os.path.join(BASE_DIR, "tickchak-orders-reserved.xlsx"))
                worksheet = workbook['כרטיסים']
                row = next([cell.value for cell in row] for row in worksheet.iter_rows(min_row=2)
                           if row[0].value == order_id)
                row[8] = ticket_desc  # כותרת
                worksheet.append(row)
                content = BytesIO()
                workbook.save(content)

                with self.assertRaisesMessage(ValueError, error):
                    _process_tickchak_orders(SimpleUploadedFile("test.xlsx", content.getvalue()), EVENT_SKU,
                                             EVENT_DATE)
                self.assertEqual(TicketOrder.objects.count(), 0)
